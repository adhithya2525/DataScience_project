from flask import Flask, request, jsonify, send_file
from flask_cors import CORS
from docx import Document
import mammoth
import requests
import mysql.connector
import os
import sys

# Set up directories
current_dir = os.path.dirname(os.path.abspath(__file__))
parent_dir = os.path.dirname(current_dir)
sys.path.append(parent_dir)

# Importing both models
from model.bot import get_response
from model.similarity import get_document

# Flask app setup
app = Flask(__name__)
CORS(app, resources={r"/api/*": {"origins": "http://localhost:3000"}})

# ✅ MySQL connection setup
db = mysql.connector.connect(
    host="localhost",       # change if your MySQL server is remote
    user="root",            # your MySQL username
    password="Jai_Sql1234",            # your MySQL password (fill this in)
    database="legal_aid_db" # your MySQL database name
)

# ---------------------- ROUTES ---------------------- #

@app.route('/api/services', methods=["GET"])
def services():
    cur = db.cursor(dictionary=True)
    cur.execute("SELECT * FROM services;")
    data = cur.fetchall()
    cur.close()
    return jsonify(data)

@app.route('/api/forms', methods=["GET"])
def get_forms():
    service_id = request.args.get('service_id')
    cur = db.cursor(dictionary=True)
    cur.execute("""
        SELECT services.service_id, services.service_name, forms.form_id, forms.form_name, forms.form_link 
        FROM services 
        INNER JOIN forms ON services.service_id = forms.service_id 
        WHERE forms.service_id = %s;
    """, (service_id,))
    data = cur.fetchall()
    cur.close()
    return jsonify(data)

@app.route('/api/form-details', methods=["GET"])
def get_form_details():
    form_id = request.args.get('form_id')
    cur = db.cursor(dictionary=True)
    cur.execute("SELECT * FROM forms WHERE form_id = %s;", (form_id,))
    form_data = cur.fetchall()

    cur.execute("""
        SELECT * FROM ques_categories 
        WHERE id IN (
            SELECT DISTINCT(category_id) 
            FROM input_ques 
            WHERE ques_id IN (
                SELECT form_query_id FROM form_queries WHERE form_id = %s
            )
        );
    """, (form_id,))
    category_data = cur.fetchall()

    cur.execute("""
        SELECT * FROM input_ques 
        WHERE ques_id IN (
            SELECT form_query_id FROM form_queries WHERE form_id = %s
        );
    """, (form_id,))
    question_data = cur.fetchall()

    cur.close()
    return jsonify({
        "form_details": form_data,
        "categories": category_data,
        "questions": question_data
    })

@app.route('/api/final-content', methods=["POST"])
def final_content():
    form_details = request.json
    form_id = form_details["form_id"]

    cur = db.cursor(dictionary=True)
    cur.execute("SELECT form_link FROM forms WHERE form_id = %s;", (form_id,))
    result = cur.fetchone()
    cur.close()

    if not result:
        return jsonify({"error": "Form not found"}), 404

    form_link = result["form_link"]
    response = requests.get(form_link)

    directory = './docs'
    os.makedirs(directory, exist_ok=True)
    file_path = './docs/localfile.docx'

    with open(file_path, 'wb') as f:
        f.write(response.content)

    doc = Document(file_path)

    # Replace placeholders like #1, #2 with form details
    replacements = {k: v for k, v in form_details.items() if k.isdigit()}
    for key, value in replacements.items():
        old = f'#{key}'
        for p in doc.paragraphs:
            if old in p.text:
                for run in p.runs:
                    if old in run.text:
                        run.text = run.text.replace(old, str(value))
    output_path = "./docs/Output2.docx"
    doc.save(output_path)

    with open(output_path, 'rb') as f:
        html_content = mammoth.convert_to_html(f)

    return jsonify({'content': html_content.value})

@app.route('/api/final-form', methods=["POST"])
def final_form():
    contents = request.get_json()
    file_path = './docs/Output2.docx'
    with open(file_path, 'w') as file:
        file.write(contents.get("content", ""))
    return send_file(file_path, as_attachment=True)

@app.route('/api/chat', methods=['POST'])
def chat():
    user_input = request.json
    user_message = user_input.get('user_chat', '')

    try:
        response = get_document(user_message)
    except:
        response = get_response(user_message)

    return jsonify({'aiMessage': response})

# ---------------------- MAIN ---------------------- #
if __name__ == '__main__':
    os.makedirs('docs', exist_ok=True)
    app.run(debug=True)
